"""docx_port.py — layered, editable DOCX export for MyPDF-OCR (wordcraft branch).

Each PDF page becomes one Word section sized like the PDF page, built from
two layers:

  1. Background layer  A picture anchored *behind* the text. Born-digital pages
                       are rendered after the PDF text has been redacted, so
                       logos, signatures, stamps and table/rule lines survive.
                       Scan-like pages are rendered and the OCR line boxes are
                       painted out with the local background colour.
  2. Text layer        One absolutely positioned Word text box per line, with
                       real, editable Thai/Latin text.

Text comes from (in order of preference): the corrected OCR text produced by
the app (Typhoon / Gemini / Tesseract, aligned line-by-line onto the page
geometry), the PDF text layer, or a Tesseract TSV pass.

Public entry point:

    docx_bytes, report = build_editable_docx(pdf_bytes, ocr_text, tesseract_cmd=...)

The Word-writing core (DocxBuilder, align_reference_text, boxes_from_tsv,
boxes_from_text_dict, erase_boxes_in_image) does not need PyMuPDF; only the
page-analysis adapter at the bottom of the file does.
"""
from __future__ import annotations

import csv
import difflib
import io
import re
import shutil
import statistics
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

DEFAULT_FONT = "TH Sarabun New"
MAX_PAGE_PT = 1584.0  # Word's page limit is 22 inches

# Accept both the app's Thai marker and the ASCII form used by older saved OCR
# results/tests. Keeping the marker parser permissive prevents the editable DOCX
# exporter from silently dropping OCR text when the label changes slightly.
PAGE_MARKER_RE = re.compile(
    r"^-----\s*(?:หน้า|page)\s+([0-9]+)(?:\s*\|.*?)?\s*-----\s*$",
    re.IGNORECASE | re.MULTILINE,
)
PLACEHOLDER_RE = re.compile(r"^\[(?:ไม่พบข้อความ|อ่านหน้านี้ไม่สำเร็จ)\]$")
THAI_RE = re.compile(r"[\u0e00-\u0e7f]")

_KNOWN_FONTS = {
    "thsarabunnew": "TH Sarabun New", "thsarabun": "TH Sarabun New",
    "sarabun": "Sarabun", "tahoma": "Tahoma", "angsananew": "Angsana New",
    "angsanaupc": "AngsanaUPC", "cordianew": "Cordia New", "cordiaupc": "CordiaUPC",
    "browalianew": "Browallia New", "browaliaupc": "BrowalliaUPC",
    "garuda": "Garuda", "norasi": "Norasi", "loma": "Loma", "leelawadee": "Leelawadee",
}


# ---------------------------------------------------------------- data model
@dataclass
class TextBox:
    text: str
    x: float
    y: float
    w: float
    h: float                 # PDF points, origin top-left
    size: float = 0.0        # 0 = estimate from geometry
    font: str = ""
    bold: bool = False
    italic: bool = False
    color: str = "000000"
    conf: float = 100.0
    inserted: bool = False
    original: str = ""


# --------------------------------------------------------------- text helpers
def normalize_pdf_text(text: str) -> str:
    """Fix Thai PDF glyph artifacts (PUA tone marks, split sara am, mark order)."""
    text = text.translate(str.maketrans({
        "\uf70a": "\u0e48", "\uf70b": "\u0e49", "\uf70e": "\u0e4c",
    }))
    text = unicodedata.normalize("NFC", text)
    text = re.sub(r"\u0e4d\s*\u0e32", "\u0e33", text)
    text = re.sub(r"([\u0e48-\u0e4b])\u0e33", "ำ\\1", text)
    text = re.sub(r"([\u0e48-\u0e4b])([\u0e38-\u0e3a])", r"\2\1", text)
    return text.strip()


def text_layer_is_corrupt(text: str) -> bool:
    if not text.strip():
        return True
    if "\ufffd" in text or any("\uf700" <= c <= "\uf8ff" for c in text):
        return True
    if re.search(r"[\x80-\x9f]|(?:Ã|Â|â|เธ)[\x80-\x9f]", text):
        return True
    thai = len(re.findall(r"[ก-๙]", text))
    if thai < 40:
        return False
    singles = len(re.findall(r"(?<![ก-๙])[ก-๙](?=\s|$)", text))
    return singles / thai > 0.08


def resolve_font(pdf_font_name: str, default: str) -> str:
    name = re.sub(r"^[A-Z]{6}\+", "", pdf_font_name or "")
    key = re.sub(r"[^a-z]", "", re.split(r"[-,]", name)[0].lower())
    return _KNOWN_FONTS.get(key, default)


def split_ocr_pages(text: str) -> dict[int, list[str]]:
    """Split the app's OCR text ('----- หน้า N | method -----' markers) per page."""
    pages: dict[int, list[str]] = {}
    if not text:
        return pages
    matches = list(PAGE_MARKER_RE.finditer(text))
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        lines = [ln.strip() for ln in text[m.end():end].splitlines()]
        pages[int(m.group(1))] = [ln for ln in lines if ln and not PLACEHOLDER_RE.match(ln)]
    return pages


def join_ocr_words(words: list[tuple[str, float, float, float]]) -> str:
    """Join OCR words (text, left, right, height). Thai has no inter-word spaces,
    so insert one only where the horizontal gap is a real word gap."""
    out = ""
    prev_right = None
    for text, left, right, height in words:
        if out:
            gap = left - prev_right
            both_thai = THAI_RE.match(out[-1]) and THAI_RE.match(text[0])
            if (not both_thai) or gap > 0.30 * max(height, 1.0):
                out += " "
        out += text
        prev_right = right
    return out


# ------------------------------------------------- geometry sources (pure)
def boxes_from_text_dict(text_dict: dict, default_font: str = DEFAULT_FONT) -> list[TextBox]:
    """PyMuPDF page.get_text('dict') → one TextBox per text line."""
    boxes: list[TextBox] = []
    for block in text_dict.get("blocks", []):
        for line in block.get("lines", []):
            spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
            if not spans:
                continue
            x0 = min(float(s["bbox"][0]) for s in spans)
            y0 = min(float(s["bbox"][1]) for s in spans)
            x1 = max(float(s["bbox"][2]) for s in spans)
            y1 = max(float(s["bbox"][3]) for s in spans)
            main = max(spans, key=lambda s: len(s.get("text", "")))
            flags = int(main.get("flags", 0))
            font_name = str(main.get("font", ""))
            color = int(main.get("color", 0))
            boxes.append(TextBox(
                text="".join(s.get("text", "") for s in spans).strip(),
                x=x0, y=y0, w=max(2.0, x1 - x0), h=max(4.0, y1 - y0),
                size=max(4.0, max(float(s.get("size", y1 - y0)) for s in spans)),
                font=resolve_font(font_name, default_font),
                bold=bool(flags & 16) or "bold" in font_name.lower(),
                italic=bool(flags & 2) or "italic" in font_name.lower() or "oblique" in font_name.lower(),
                color=f"{(color >> 16) & 255:02X}{(color >> 8) & 255:02X}{color & 255:02X}",
            ))
    return boxes


def boxes_from_tsv(tsv_text: str, scale: float, min_word_conf: float = 30.0,
                   min_line_conf: float = 35.0) -> list[TextBox]:
    """Tesseract TSV → one TextBox per line (scale converts pixels → points)."""
    rows = csv.DictReader(io.StringIO(tsv_text), delimiter="\t", quoting=csv.QUOTE_NONE)
    lines: dict[tuple, list[dict]] = {}
    for row in rows:
        text = (row.get("text") or "").strip()
        try:
            conf = float(row.get("conf", "-1"))
            geom = [int(row[k]) for k in ("left", "top", "width", "height")]
        except (ValueError, KeyError, TypeError):
            continue
        if not text or conf < min_word_conf:
            continue
        key = (row.get("page_num"), row.get("block_num"), row.get("par_num"), row.get("line_num"))
        lines.setdefault(key, []).append({"text": text, "conf": conf,
                                          "l": geom[0], "t": geom[1], "w": geom[2], "h": geom[3]})
    boxes: list[TextBox] = []
    for words in lines.values():
        words.sort(key=lambda d: d["l"])
        mean_conf = sum(d["conf"] for d in words) / len(words)
        left = min(d["l"] for d in words)
        top = min(d["t"] for d in words)
        right = max(d["l"] + d["w"] for d in words)
        bottom = max(d["t"] + d["h"] for d in words)
        height = bottom - top
        text = join_ocr_words([(d["text"], d["l"], d["l"] + d["w"], height) for d in words])
        text = normalize_pdf_text(text)
        if mean_conf < min_line_conf or len(text.replace(" ", "")) < 2:
            continue
        boxes.append(TextBox(text=text, x=left * scale, y=top * scale,
                             w=(right - left) * scale, h=height * scale, conf=mean_conf))
    boxes.sort(key=lambda b: (round(b.y / 4.0), b.x))
    return boxes


# --------------------------------------------------- reference-text alignment
def _norm(s: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", s))


def _norm_with_map(s: str) -> tuple[str, list[int], str]:
    """Whitespace-free NFC text, the index of each kept char in the NFC string, the NFC string."""
    nfc = unicodedata.normalize("NFC", s)
    chars, idx = [], []
    for i, ch in enumerate(nfc):
        if not ch.isspace():
            chars.append(ch)
            idx.append(i)
    return "".join(chars), idx, nfc


def _ratio(a: str, b: str, threshold: float) -> float:
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    if sm.real_quick_ratio() < threshold or sm.quick_ratio() < threshold:
        return 0.0
    r = sm.ratio()
    return r if r >= threshold else 0.0


def _rescue_windows(boxes, nb, ref_lines, unmatched_boxes, unmatched_refs, threshold):
    """Order-independent pass for table rows / merged lines: a box takes the
    best-matching *substring* of a reference line that no other box has taken."""
    refs = {j: _norm_with_map(ref_lines[j]) for j in unmatched_refs}
    cands = []
    for bi in unmatched_boxes:
        a = nb[bi]
        L = len(a)
        if L < 2:
            continue
        for j, (rn, _idx, _nfc) in refs.items():
            R = len(rn)
            if R < L * 0.4:
                continue
            seen = set()
            for f in (0.6, 0.8, 1.0, 1.2):
                ln = min(R, max(2, round(L * f)))
                if ln in seen:
                    continue
                seen.add(ln)
                for s in range(0, R - ln + 1):
                    r = _ratio(a, rn[s:s + ln], threshold)
                    if r:
                        cands.append((r, bi, j, s, s + ln))
    cands.sort(key=lambda c: -c[0])
    used: dict[int, list[tuple[int, int]]] = {}
    done: dict[int, tuple[int, str]] = {}
    for r, bi, j, s, e in cands:
        if bi in done:
            continue
        if any(min(e, ue) - max(s, us) > 0.3 * (e - s) for us, ue in used.get(j, [])):
            continue
        rn, idx, nfc = refs[j]
        used.setdefault(j, []).append((s, e))
        done[bi] = (j, nfc[idx[s]: idx[e - 1] + 1])
    for bi, (_j, text) in done.items():
        boxes[bi].original = boxes[bi].text
        boxes[bi].text = text
    return done, used


def align_reference_text(boxes: list[TextBox], ref_lines: list[str],
                         threshold: float = 0.55, rescue_threshold: float = 0.6) -> dict:
    """Replace box text by the corrected reference lines.

    Pass 1: monotone one-to-one line alignment.
    Pass 2: substring rescue for boxes left over (table cells, merged lines).
    Boxes that still match nothing keep their own text. Reference lines that are
    not represented anywhere are inserted under the previous matched box.
    """
    n, m = len(boxes), len(ref_lines)
    stats = {"boxes": n, "ref_lines": m, "matched": 0, "rescued": 0, "kept": n, "inserted": 0}
    if not n or not m:
        return stats
    nb = [_norm(b.text) for b in boxes]
    nr = [_norm(t) for t in ref_lines]
    window = max(12, abs(n - m) + 8)
    sim = [[0.0] * m for _ in range(n)]
    for i in range(n):
        centre = i * m / n
        for j in range(m):
            if abs(j - centre) > window or not nb[i] or not nr[j]:
                continue
            # A corrected line may be shorter than a garbled box (extra marks), but
            # must not be much longer than the box, or it would overflow it.
            if not 0.4 <= len(nr[j]) / len(nb[i]) <= 1.25:
                continue
            sim[i][j] = _ratio(nb[i], nr[j], threshold)
    dp = [[0.0] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            best = max(dp[i - 1][j], dp[i][j - 1])
            if sim[i - 1][j - 1] > 0:
                best = max(best, dp[i - 1][j - 1] + sim[i - 1][j - 1])
            dp[i][j] = best
    pairs: dict[int, int] = {}
    i, j = n, m
    while i > 0 and j > 0:
        if sim[i - 1][j - 1] > 0 and abs(dp[i][j] - dp[i - 1][j - 1] - sim[i - 1][j - 1]) < 1e-9:
            pairs[i - 1] = j - 1
            i -= 1
            j -= 1
        elif abs(dp[i][j] - dp[i - 1][j]) < 1e-9:
            i -= 1
        else:
            j -= 1
    for bi, rj in pairs.items():
        boxes[bi].original = boxes[bi].text
        boxes[bi].text = ref_lines[rj]
    stats["matched"] = len(pairs)

    matched_refs = set(pairs.values())
    un_boxes = [i for i in range(n) if i not in pairs]
    un_refs = [j for j in range(m) if j not in matched_refs]
    rescued, used = _rescue_windows(boxes, nb, ref_lines, un_boxes, un_refs, rescue_threshold)
    stats["rescued"] = len(rescued)
    stats["kept"] = n - len(pairs) - len(rescued)

    # Reference lines no box represents: insert them (skip when mostly covered).
    leftover = "".join(nb[i] for i in un_boxes if i not in rescued)
    inserts: list[tuple[int, int]] = []   # (after box index, ref index)
    for rj in un_refs:
        if len(nr[rj]) < 3:
            continue
        covered_chars = sum(e - s for s, e in used.get(rj, []))
        if covered_chars / len(nr[rj]) >= 0.6:
            continue
        if leftover:
            sm = difflib.SequenceMatcher(None, nr[rj], leftover, autojunk=False)
            if sum(b.size for b in sm.get_matching_blocks()) / len(nr[rj]) >= 0.6:
                continue
        prev_boxes = [bi for bi, rr in pairs.items() if rr < rj]
        inserts.append((max(prev_boxes) if prev_boxes else -1, rj))
    extra: list[TextBox] = []
    offsets: dict[int, int] = {}
    for after, rj in inserts:
        anchor = boxes[after] if after >= 0 else boxes[0]
        k = offsets.get(after, 0) + 1
        offsets[after] = k
        extra.append(TextBox(
            text=ref_lines[rj], x=anchor.x,
            y=anchor.y + anchor.h * 1.2 * k if after >= 0 else max(0.0, anchor.y - anchor.h * 1.2 * k),
            w=anchor.w, h=anchor.h, size=anchor.size, font=anchor.font,
            bold=anchor.bold, italic=anchor.italic, color=anchor.color, inserted=True))
    boxes.extend(extra)
    stats["inserted"] = len(extra)
    return stats


# ------------------------------------------------------ font measuring / sizing
def find_font_file(font_dir, family: str):
    """Locate a regular-weight TTF/OTF for `family` to measure text widths."""
    if not font_dir or not Path(font_dir).is_dir():
        return None
    files = [p for p in Path(font_dir).rglob("*") if p.suffix.lower() in (".ttf", ".otf")]
    want = re.sub(r"[^a-z]", "", family.lower())
    for needle in (want, "sarabun"):
        for p in files:
            stem = re.sub(r"[^a-z]", "", p.stem.lower())
            if needle in stem and not any(k in stem for k in ("bold", "italic", "oblique")):
                return p
    return None


class TextMeasurer:
    def __init__(self, font_path):
        from PIL import ImageFont
        self._font = ImageFont.truetype(str(font_path), 100)

    def length(self, text: str, size: float) -> float:
        return float(self._font.getlength(text)) * size / 100.0


def resolve_sizes(boxes: list[TextBox], measurer) -> list[float]:
    """Font size per box. Explicit sizes win; otherwise fit the text width to the
    line box (when the font can be measured) and snap near-equal sizes together."""
    sizes: list[float] = []
    fitted: list[int] = []
    for idx, b in enumerate(boxes):
        if b.size > 0:
            sizes.append(b.size)
            continue
        base = max(4.0, b.h * 0.72)
        fs = base
        if measurer is not None and len(b.text.strip()) >= 4 and b.w > 0:
            length = measurer.length(b.text, 100.0)
            if length > 0:
                fs = min(max(b.w / length * 100.0, base * 0.5), base * 2.2)
                if len(b.text.strip()) >= 10:
                    fitted.append(idx)
        sizes.append(fs)
    if len(fitted) >= 3:
        median = statistics.median(sizes[i] for i in fitted)
        for i, b in enumerate(boxes):
            if b.size <= 0 and abs(sizes[i] - median) <= 0.12 * median:
                sizes[i] = median
    return [max(4.0, s) for s in sizes]


# ------------------------------------------------------------- image helpers
def erase_boxes_in_image(img, boxes: list[TextBox], scale: float, pad_pt: float = 1.5):
    """Paint boxes out with the median colour of a thin ring around each box."""
    from PIL import ImageDraw, ImageStat
    img = img.convert("RGB")
    draw = ImageDraw.Draw(img)
    W, H = img.size
    for b in boxes:
        pad_x = pad_pt * scale
        pad_top = max(pad_pt, 0.35 * b.h) * scale      # Thai tone marks sit above the line
        pad_bottom = max(pad_pt, 0.2 * b.h) * scale
        x0 = max(0, int(b.x * scale - pad_x))
        y0 = max(0, int(b.y * scale - pad_top))
        x1 = min(W, int((b.x + b.w) * scale + pad_x))
        y1 = min(H, int((b.y + b.h) * scale + pad_bottom))
        if x1 - x0 < 2 or y1 - y0 < 2:
            continue
        ring = []
        for strip in ((x0, max(0, y0 - 4), x1, y0), (x0, y1, x1, min(H, y1 + 4))):
            if strip[2] > strip[0] and strip[3] > strip[1]:
                ring.append(ImageStat.Stat(img.crop(strip)).median)
        if ring:
            fill = tuple(int(sum(c[k] for c in ring) / len(ring)) for k in range(3))
        else:
            fill = (255, 255, 255)
        draw.rectangle([x0, y0, x1, y1], fill=fill)
    return img


# ------------------------------------------------------------------ DOCX core
_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
       'xmlns:v="urn:schemas-microsoft-com:vml" '
       'xmlns:o="urn:schemas-microsoft-com:office:office"')


def _inline_to_anchor_behind(run) -> None:
    """Turn the inline picture python-docx just made into a page-anchored picture
    placed behind the text."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    inline = run._r.find(".//" + qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{extent.get("cx")}" cy="{extent.get("cy")}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.insert(list(anchor).index(anchor.find(qn("wp:cNvGraphicFramePr"))), doc_pr)
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


class DocxBuilder:
    def __init__(self, default_font: str = DEFAULT_FONT, measurer=None, y_adjust: float = 0.0):
        from docx import Document
        self.default_font = default_font
        self.measurer = measurer
        self.y_adjust = y_adjust
        self.doc = Document()
        zoom = self.doc.settings.element.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}zoom")
        if zoom is not None and zoom.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent") is None:
            zoom.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent", "100")
        self._first = True
        self._shape_id = 0
        normal = self.doc.styles["Normal"]
        normal.font.name = default_font
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
        if fonts is not None:
            for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                fonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}", default_font)

    # -- page ---------------------------------------------------------------
    def add_page(self, width_pt: float, height_pt: float, boxes: list[TextBox],
                 background: tuple[bytes, str] | None = None,
                 flow_lines: list[str] | None = None) -> None:
        from docx.enum.section import WD_SECTION
        from docx.shared import Pt, Inches
        k = min(1.0, MAX_PAGE_PT / max(width_pt, height_pt, 1.0))
        section = self.doc.sections[0] if self._first else self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._first = False
        section.page_width, section.page_height = Pt(width_pt * k), Pt(height_pt * k)
        margin = Inches(0.8) if flow_lines else Inches(0)
        section.top_margin = section.bottom_margin = margin
        section.left_margin = section.right_margin = margin
        section.header_distance = section.footer_distance = Inches(0)

        paragraph = self.doc.add_paragraph()
        pf = paragraph.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = Pt(1)
        if background is not None:
            data, _ext = background
            run = paragraph.add_run()
            run.add_picture(io.BytesIO(data), width=Pt(width_pt * k), height=Pt(height_pt * k))
            _inline_to_anchor_behind(run)

        # If page text has no usable coordinates (for example a scanned PDF
        # without local Tesseract language data), keep the OCR result instead
        # of silently exporting only the page image. The paragraphs are real
        # editable Word text and the PDF image remains behind them.
        if flow_lines:
            for line in flow_lines:
                text_paragraph = self.doc.add_paragraph(line)
                text_paragraph.paragraph_format.space_before = Pt(0)
                text_paragraph.paragraph_format.space_after = Pt(0)
            return

        sizes = resolve_sizes(boxes, self.measurer)
        for z, (box, size) in enumerate(zip(boxes, sizes), start=1):
            self._add_box(paragraph, box, size * k, k, z)

    # -- text box -----------------------------------------------------------
    def _add_box(self, paragraph, box: TextBox, size: float, k: float, z: int) -> None:
        from docx.oxml import parse_xml
        text = box.text.strip()
        if not text:
            return
        scale_w = 100
        has_thai = bool(THAI_RE.search(text))
        if self.measurer is not None:
            measured = self.measurer.length(text, size)
            if box.original and measured > box.w * k * 1.4:
                if not text_layer_is_corrupt(box.original):
                    text = box.original.strip() or text      # only revert if original text was clean
                    measured = self.measurer.length(text, size)
            if measured > box.w * k * 1.02 > 0:
                scale_w = int(min(100, max(50, round(box.w * k / measured * 100))))
        shape_h = size * (1.85 if has_thai else 1.5)
        cy = (box.y + box.h / 2.0) * k
        top = cy - shape_h / 2.0 + self.y_adjust * size
        left, width = box.x * k, box.w * k * 1.04 + 3.0
        self._shape_id += 1
        font = escape(box.font or self.default_font)
        rpr = (f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:eastAsia="{font}"/>'
               + ('<w:b/><w:bCs/>' if box.bold else '')
               + ('<w:i/><w:iCs/>' if box.italic else '')
               + f'<w:color w:val="{box.color}"/>'
               + (f'<w:w w:val="{scale_w}"/>' if scale_w != 100 else '')
               + f'<w:sz w:val="{max(2, round(size * 2))}"/><w:szCs w:val="{max(2, round(size * 2))}"/>'
               + '<w:lang w:val="th-TH" w:bidi="th-TH"/>')
        xml = (
            f'<w:r {_NS}><w:pict>'
            f'<v:rect id="pdfbox{self._shape_id}" stroked="f" filled="f" style="position:absolute;'
            f'margin-left:{left:.2f}pt;margin-top:{top:.2f}pt;width:{width:.2f}pt;height:{shape_h:.2f}pt;'
            f'z-index:{z};mso-position-horizontal-relative:page;mso-position-vertical-relative:page;'
            'mso-wrap-style:none;v-text-anchor:top">'
            '<v:textbox inset="0,0,0,0"><w:txbxContent><w:p><w:pPr>'
            '<w:wordWrap w:val="0"/><w:snapToGrid w:val="0"/>'
            f'<w:spacing w:before="0" w:after="0" w:line="{round(shape_h * 20)}" w:lineRule="atLeast"/>'
            f'</w:pPr><w:r><w:rPr>{rpr}</w:rPr>'
            f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p></w:txbxContent></v:textbox>'
            '</v:rect></w:pict></w:r>')
        paragraph._p.append(parse_xml(xml))

    def save(self, title: str = "MyOCR PDF editable DOCX export") -> bytes:
        cp = self.doc.core_properties
        cp.title = title
        cp.subject = "PDF rebuilt as background layer + editable Word text"
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()


# ============================================================================
# PyMuPDF adapter (needs `pymupdf`, Pillow; Tesseract optional)
# ============================================================================
def _import_fitz():
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    return fitz


def find_tesseract(explicit: str | None = None) -> str | None:
    if explicit:
        return explicit
    found = shutil.which("tesseract")
    if found:
        return found
    for cand in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                 r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
        if Path(cand).exists():
            return cand
    return None


_THAI_OK: dict[tuple, bool] = {}


def _tesseract_has_thai(command: str, tessdata_dir) -> bool:
    """Tesseract silently falls back to English when `tha` is missing, which yields
    garbage for Thai pages, so check before using it."""
    key = (command, str(tessdata_dir))
    if key not in _THAI_OK:
        args = [command, "--list-langs"]
        if tessdata_dir and (Path(tessdata_dir) / "tha.traineddata").exists():
            args += ["--tessdata-dir", str(tessdata_dir)]
        try:
            out = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30, check=False)
            _THAI_OK[key] = "tha" in (out.stdout + out.stderr).decode("utf-8", errors="replace").split()
        except Exception:
            _THAI_OK[key] = False
    return _THAI_OK[key]


def _ocr_boxes(page, command: str | None, tessdata_dir) -> list[TextBox]:
    if not command or not _tesseract_has_thai(command, tessdata_dir):
        return []
    dpi = 300
    png = page.get_pixmap(dpi=dpi, alpha=False).tobytes("png")
    with tempfile.TemporaryDirectory(prefix="docx_port_") as tmp:
        image_path = Path(tmp) / "page.png"
        image_path.write_bytes(png)
        args = [command, str(image_path), "stdout", "-l", "tha+eng", "--psm", "3",
                "--oem", "1", "--dpi", str(dpi)]
        if tessdata_dir and (Path(tessdata_dir) / "tha.traineddata").exists():
            args += ["--tessdata-dir", str(tessdata_dir)]
        args.append("tsv")
        result = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                timeout=180, check=False)
    if result.returncode != 0:
        return []
    return boxes_from_tsv(result.stdout.decode("utf-8", errors="replace"), 72.0 / dpi)


def _is_scan_like(page, fitz) -> bool:
    """A page dominated by one image is a scan (even if it has a hidden OCR layer)."""
    area = max(1.0, float(page.rect.width * page.rect.height))
    try:
        for img in page.get_images(full=True):
            for rect in page.get_image_rects(img[0]):
                if float(rect.width * rect.height) / area > 0.6:
                    return True
    except Exception:
        return False
    return False


def _render_background(scratch, index: int, text_rects: list, placed: list[TextBox],
                       scan_like: bool, dpi: int, fitz) -> tuple[tuple[bytes, str], str]:
    from PIL import Image
    page = scratch[index]
    if scan_like:
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        img = erase_boxes_in_image(img, placed, dpi / 72.0)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=82, optimize=True)
        return (buf.getvalue(), "jpeg"), "erased"
    status = "plain"
    if text_rects:
        for r in text_rects:
            rect = fitz.Rect(r) + (-0.5, -0.5, 0.5, 0.5)
            if page.rotation:
                rect = rect * page.derotation_matrix
            page.add_redact_annot(rect, fill=False)
        none_img = getattr(fitz, "PDF_REDACT_IMAGE_NONE", 0)
        none_art = getattr(fitz, "PDF_REDACT_LINE_ART_NONE", 0)
        try:
            page.apply_redactions(images=none_img, graphics=none_art)
        except TypeError:
            page.apply_redactions(images=none_img)
        status = "redacted"
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    return (pix.tobytes("png"), "png"), status


def build_editable_docx(pdf_bytes: bytes, ocr_text: str = "", *,
                        tesseract_cmd: str | None = None, tessdata_dir=None,
                        font_dir=None, default_font: str = DEFAULT_FONT,
                        keep_background: bool = True, prefer_text_layer_for_thai: bool = False,
                        dpi: int = 200, y_adjust: float = 0.0) -> tuple[bytes, dict]:
    """Rebuild a PDF as an editable DOCX. Returns (docx_bytes, report)."""
    if not pdf_bytes:
        raise ValueError("ไม่พบข้อมูล PDF สำหรับสร้าง DOCX")
    fitz = _import_fitz()
    command = find_tesseract(tesseract_cmd)
    font_file = find_font_file(font_dir, default_font)
    measurer = None
    if font_file is not None:
        try:
            measurer = TextMeasurer(font_file)
        except Exception:
            measurer = None
    ref_pages = split_ocr_pages(ocr_text)
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    scratch = fitz.open(stream=pdf_bytes, filetype="pdf")
    builder = DocxBuilder(default_font, measurer, y_adjust)
    report: dict = {"pages": [], "measured_font": str(font_file) if font_file else None,
                    "tesseract": bool(command), "warnings": []}
    try:
        if doc.page_count == 0:
            raise ValueError("ไม่พบหน้าเอกสารใน PDF")
        for index in range(doc.page_count):
            page = doc[index]
            page_no = index + 1
            info: dict = {"page": page_no}
            t_boxes = boxes_from_text_dict(page.get_text("dict"), default_font)
            for b in t_boxes:
                b.text = normalize_pdf_text(b.text)
            t_boxes = [b for b in t_boxes if b.text]
            raw = "\n".join(b.text for b in t_boxes)
            thai = len(re.findall(r"[ก-๙]", raw))
            corrupt = text_layer_is_corrupt(raw)
            ref = ref_pages.get(page_no) or []

            boxes: list[TextBox] = []
            source = "none"
            if t_boxes and not corrupt:
                if ref:
                    boxes, source = t_boxes, "text-layer+reference"
                elif thai >= 20 and not prefer_text_layer_for_thai:
                    boxes = _ocr_boxes(page, command, tessdata_dir)
                    source = "ocr"
                    if not boxes:
                        boxes, source = t_boxes, "text-layer"
                else:
                    boxes, source = t_boxes, "text-layer"
            else:
                boxes = _ocr_boxes(page, command, tessdata_dir)
                source = "ocr+reference" if (boxes and ref) else ("ocr" if boxes else "none")
                if not boxes and t_boxes:
                    boxes, source = t_boxes, "text-layer(corrupt)"
                    report["warnings"].append(f"หน้า {page_no}: text layer เสียและไม่มี Tesseract — ข้อความอาจผิด")
            if ref and boxes:
                info.update(align_reference_text(boxes, ref))
            info["source"] = source

            rect = page.rect
            if not boxes and ref:
                builder.add_page(rect.width, rect.height, [], flow_lines=ref)
                info["layout"] = "flow-fallback"
                report["warnings"].append(f"หน้า {page_no}: ไม่มีพิกัดข้อความ จึงวางเป็นย่อหน้าธรรมดา (ไม่คงเลย์เอาต์)")
            else:
                scan_like = _is_scan_like(page, fitz)
                background = None
                if keep_background:
                    # Only erase/redact when we actually replace that text.
                    rects = [fitz.Rect(b.x, b.y, b.x + b.w, b.y + b.h) for b in t_boxes] if not scan_like else []
                    background, status = _render_background(
                        scratch, index, rects, [b for b in boxes if not b.inserted], scan_like, dpi, fitz)
                    info["background"] = status
                builder.add_page(rect.width, rect.height, boxes, background)
                info["layout"] = "scan" if scan_like else "born-digital"
            info["text_boxes"] = len(boxes)
            report["pages"].append(info)
        return builder.save(), report
    finally:
        doc.close()
        scratch.close()
